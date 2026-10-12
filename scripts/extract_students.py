import openpyxl
import json
import os

excel_path = r'C:/Users/USER/.gemini/antigravity/brain/5cb16225-4200-48c0-881a-e77ca1765b51/.user_uploaded/media_1791727816306_070e7c6f.xlsx'
wb = openpyxl.load_workbook(excel_path)
students = []

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    class_grade = sheet_name.replace('KELAS ', '').strip()
    session = 'Sesi 1' if class_grade.startswith('7') else 'Sesi 2'
    class_list = []
    
    for r in range(5, sheet.max_row + 1):
        name_val = sheet.cell(r, 2).value
        if name_val and str(name_val).strip() and not str(name_val).strip().lower().startswith(('rata', 'jumlah', 'keterangan')):
            no_val = sheet.cell(r, 1).value
            try:
                num = int(float(no_val)) if no_val is not None else len(class_list) + 1
            except Exception:
                num = len(class_list) + 1
            name = str(name_val).strip()
            item = {
                'invitation_code': f'ORTU-{class_grade}-{num:02d}',
                'no': num,
                'name': name,
                'class_grade': class_grade,
                'type': 'parent',
                'session': session,
                'original_session': session,
                'ticket_qty': 2,
                'is_swapped': 0,
                'swap_notes': '',
                'phone': '',
                'checked_in': 0,
                'checked_in_at': None
            }
            class_list.append(item)
            
    print(f"{class_grade}: {len(class_list)} students. First: {class_list[0]['name']}, Last: {class_list[-1]['name']}")
    students.extend(class_list)

print(f"Total students extracted: {len(students)}")

os.makedirs('config', exist_ok=True)
out_path = os.path.join('config', 'students_data.json')
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(students, f, indent=2, ensure_ascii=False)

print(f"Successfully saved {len(students)} students to {out_path}")
